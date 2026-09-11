"""
resume_utils.py
- Extract raw text from uploaded PDF / DOCX / TXT resumes.
- Data model + helpers for the manual "Enter Information Manually" form.
- Convert structured resume data into a single plain-text representation
  used for AI prompts and ATS-style resume generation.
"""

import io
import logging

logger = logging.getLogger("applyready.resume")


def empty_resume_data():
    """Return a blank resume data structure for the manual entry form."""
    return {
        "personal": {
            "full_name": "",
            "email": "",
            "phone": "",
            "location": "",
            "linkedin": "",
            "portfolio": "",
            "title": "",
        },
        "summary": "",
        "education": [],      # list of dicts
        "experience": [],     # list of dicts
        "internships": [],    # list of dicts
        "projects": [],       # list of dicts
        "skills": {
            "technical": "",
            "soft": "",
            "tools": "",
            "languages": "",
        },
        "certifications": "",
        "awards": "",
        "volunteer": "",
        "references": "",
    }


def empty_education_entry():
    return {
        "degree": "", "institution": "", "location": "",
        "start_year": "", "end_year": "", "gpa": "", "coursework": "",
    }


def empty_experience_entry():
    return {
        "title": "", "company": "", "location": "",
        "start_date": "", "end_date": "", "responsibilities": "", "achievements": "",
    }


def empty_project_entry():
    return {
        "name": "", "description": "", "technologies": "", "role": "", "achievements": "",
    }


# ---------------------------------------------------------------------
# File text extraction
# ---------------------------------------------------------------------

def extract_text_from_upload(uploaded_file):
    """
    Extract text from an uploaded file (Streamlit UploadedFile).
    Supports .pdf, .docx, .txt.
    Returns (text, error). If error is not None, text will be "" and the
    caller should show a friendly message rather than crashing.
    """
    if uploaded_file is None:
        return "", "No file provided."

    name = (uploaded_file.name or "").lower()
    raw_bytes = uploaded_file.read()
    if not raw_bytes:
        return "", "The uploaded file appears to be empty."

    try:
        if name.endswith(".pdf"):
            return _extract_pdf_text(raw_bytes)
        elif name.endswith(".docx"):
            return _extract_docx_text(raw_bytes)
        elif name.endswith(".txt"):
            return _extract_txt_text(raw_bytes)
        else:
            return "", "Unsupported file type. Please upload a PDF, DOCX, or TXT file."
    except Exception as e:
        logger.exception("Text extraction failed for %s", name)
        return "", f"Could not read this file ({e}). Try re-saving it or use manual entry instead."


def _extract_pdf_text(raw_bytes):
    text = ""
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(raw_bytes))
        pages_text = []
        for page in reader.pages:
            try:
                pages_text.append(page.extract_text() or "")
            except Exception:
                continue
        text = "\n".join(pages_text).strip()
    except Exception:
        # Fall back to PyMuPDF if pypdf isn't available or fails
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(stream=raw_bytes, filetype="pdf")
            pages_text = [p.get_text() for p in doc]
            text = "\n".join(pages_text).strip()
        except Exception as e:
            return "", f"Could not extract text from this PDF ({e}). It may be a scanned image."

    if not text:
        return "", ("No readable text was found in this PDF. It may be a scanned image "
                     "without embedded text. Please try manual entry instead.")
    return text, None


def _extract_docx_text(raw_bytes):
    try:
        import docx
        document = docx.Document(io.BytesIO(raw_bytes))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        text = "\n".join(t for t in parts if t and t.strip()).strip()
    except Exception as e:
        return "", f"Could not read this DOCX file ({e})."

    if not text:
        return "", "This DOCX file appears to contain no readable text."
    return text, None


def _extract_txt_text(raw_bytes):
    for encoding in ("utf-8", "latin-1"):
        try:
            text = raw_bytes.decode(encoding).strip()
            if text:
                return text, None
        except Exception:
            continue
    return "", "Could not decode this text file."


# ---------------------------------------------------------------------
# Structured data -> plain text
# ---------------------------------------------------------------------

def resume_data_to_text(data):
    """
    Convert the structured manual-entry resume_data dict into a clean
    plain-text resume representation, suitable for AI prompts and as a
    base for document generation.
    """
    if not data:
        return ""

    lines = []
    p = data.get("personal", {})
    name_line = p.get("full_name", "").strip()
    if name_line:
        lines.append(name_line.upper())
    if p.get("title"):
        lines.append(p["title"])

    contact_bits = [b for b in [p.get("email"), p.get("phone"), p.get("location")] if b]
    if contact_bits:
        lines.append(" | ".join(contact_bits))
    link_bits = [b for b in [p.get("linkedin"), p.get("portfolio")] if b]
    if link_bits:
        lines.append(" | ".join(link_bits))

    if data.get("summary"):
        lines.append("\nPROFESSIONAL SUMMARY")
        lines.append(data["summary"].strip())

    if data.get("experience"):
        lines.append("\nEXPERIENCE")
        for e in data["experience"]:
            if not any(e.values()):
                continue
            header = " - ".join(x for x in [e.get("title"), e.get("company")] if x)
            meta = " | ".join(x for x in [e.get("location"), f"{e.get('start_date','')} - {e.get('end_date','')}"] if x)
            if header:
                lines.append(f"{header}")
            if meta.strip(" |-"):
                lines.append(meta)
            if e.get("responsibilities"):
                for r in _split_lines(e["responsibilities"]):
                    lines.append(f"- {r}")
            if e.get("achievements"):
                for a in _split_lines(e["achievements"]):
                    lines.append(f"- {a}")

    if data.get("internships"):
        real = [i for i in data["internships"] if any(i.values())]
        if real:
            lines.append("\nINTERNSHIPS")
            for e in real:
                header = " - ".join(x for x in [e.get("title"), e.get("company")] if x)
                if header:
                    lines.append(header)
                if e.get("responsibilities"):
                    for r in _split_lines(e["responsibilities"]):
                        lines.append(f"- {r}")

    if data.get("education"):
        real = [ed for ed in data["education"] if any(ed.values())]
        if real:
            lines.append("\nEDUCATION")
            for ed in real:
                header = " - ".join(x for x in [ed.get("degree"), ed.get("institution")] if x)
                meta = " | ".join(x for x in [ed.get("location"), f"{ed.get('start_year','')} - {ed.get('end_year','')}"] if x)
                if header:
                    lines.append(header)
                if meta.strip(" |-"):
                    lines.append(meta)
                if ed.get("gpa"):
                    lines.append(f"GPA: {ed['gpa']}")
                if ed.get("coursework"):
                    lines.append(f"Relevant coursework: {ed['coursework']}")

    if data.get("projects"):
        real = [pr for pr in data["projects"] if any(pr.values())]
        if real:
            lines.append("\nPROJECTS")
            for pr in real:
                header = " - ".join(x for x in [pr.get("name"), pr.get("role")] if x)
                if header:
                    lines.append(header)
                if pr.get("description"):
                    lines.append(pr["description"])
                if pr.get("technologies"):
                    lines.append(f"Technologies: {pr['technologies']}")
                if pr.get("achievements"):
                    for a in _split_lines(pr["achievements"]):
                        lines.append(f"- {a}")

    skills = data.get("skills", {})
    if any(skills.values()):
        lines.append("\nSKILLS")
        if skills.get("technical"):
            lines.append(f"Technical: {skills['technical']}")
        if skills.get("tools"):
            lines.append(f"Tools: {skills['tools']}")
        if skills.get("soft"):
            lines.append(f"Soft Skills: {skills['soft']}")
        if skills.get("languages"):
            lines.append(f"Languages: {skills['languages']}")

    if data.get("certifications"):
        lines.append("\nCERTIFICATIONS")
        lines.append(data["certifications"].strip())

    if data.get("awards"):
        lines.append("\nAWARDS / ACHIEVEMENTS")
        lines.append(data["awards"].strip())

    if data.get("volunteer"):
        lines.append("\nVOLUNTEER / EXTRACURRICULAR")
        lines.append(data["volunteer"].strip())

    if data.get("references"):
        lines.append("\nREFERENCES")
        lines.append(data["references"].strip())

    return "\n".join(lines).strip()


def _split_lines(block):
    """Split a multi-line textarea entry into clean non-empty lines."""
    if not block:
        return []
    out = []
    for line in block.split("\n"):
        line = line.strip("-• \t")
        if line:
            out.append(line)
    return out
